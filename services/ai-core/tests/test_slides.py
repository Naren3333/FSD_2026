"""PowerPoint decks: the deck reader, picture descriptions and POST /v1/ingest with a .pptx."""

import io
import json
import zipfile
from pathlib import Path
from typing import Any

import pymupdf
import pytest
from conftest import FakeOpenAI, chat_response, fake_embeddings, ingested, with_vision
from fastapi.testclient import TestClient
from lxml import etree  # type: ignore[import-untyped]
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from ingest import PPTX, VisionConfig, slides
from ingest.chunker import DESCRIBED
from ingest.parse import IngestError

DESCRIPTION = "A pipeline: commit, build, test and deploy, each feeding the next."


def picture(width: int = 400, height: int = 300, shade: int = 90) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), False)
    pix.set_rect(pix.irect, (shade, 120, 200))
    return pix.tobytes("png")


def make_deck(*, sections: bool = False, hidden: bool = False, logo_everywhere: bool = False) -> bytes:
    prs = Presentation()
    title_only, title_body, blank = prs.slide_layouts[5], prs.slide_layouts[1], prs.slide_layouts[6]

    s1 = prs.slides.add_slide(title_body)
    s1.shapes.title.text = "Continuous Integration"
    body = s1.placeholders[1].text_frame
    body.text = "Merge small changes often"
    sub = body.add_paragraph()
    sub.text, sub.level = "Every push is built", 1
    s1.notes_slide.notes_text_frame.text = "Ask who has broken the build this week."

    s2 = prs.slides.add_slide(title_only)
    s2.shapes.title.text = "Build times"
    rows = s2.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, cells in enumerate([("Stage", "Minutes"), ("Build", "4"), ("Test", "11")]):
        for c, text in enumerate(cells):
            rows.cell(r, c).text = text
    data = CategoryChartData()
    data.categories = ["Mon", "Tue"]
    data.add_series("Builds", (12, 7.5))
    chart = s2.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4.5), Inches(4), Inches(2.5), data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Builds per day"

    s3 = prs.slides.add_slide(blank)  # no title: stays under "Build times"
    s3.shapes.add_picture(io.BytesIO(picture()), Inches(1), Inches(1), Inches(6), Inches(4))
    s3.shapes.add_picture(io.BytesIO(picture(20, 20, 10)), Inches(9), Inches(7), Inches(0.3), Inches(0.3))
    box = s3.shapes.add_textbox(Inches(1), Inches(5.5), Inches(6), Inches(1)).text_frame
    box.text = "The pipeline runs on every push."

    s4 = prs.slides.add_slide(title_body)
    s4.shapes.title.text = "Secret slide"
    s4.placeholders[1].text_frame.text = "Never shown"
    if hidden:
        s4._element.set("show", "0")

    if logo_everywhere:
        for s in prs.slides:
            s.shapes.add_picture(
                io.BytesIO(picture(300, 300, 30)), Inches(7), Inches(5), Inches(2), Inches(2)
            )

    out = io.BytesIO()
    prs.save(out)
    data_ = out.getvalue()
    return _with_sections(data_, prs) if sections else data_


def _with_sections(data: bytes, prs: Any) -> bytes:
    """PowerPoint's sections (python-pptx can't write them): "Basics" holds slide 1, "Pipelines" the rest."""
    ids = [s.slide_id for s in prs.slides]
    p14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            blob = src.read(item.filename)
            if item.filename == "ppt/presentation.xml":
                root = etree.fromstring(blob)
                ext_lst = etree.SubElement(root, f"{{{p}}}extLst")
                ext = etree.SubElement(ext_lst, f"{{{p}}}ext", uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
                section_lst = etree.SubElement(ext, f"{{{p14}}}sectionLst")
                for name, members in [("Basics", ids[:1]), ("Pipelines", ids[1:])]:
                    section = etree.SubElement(section_lst, f"{{{p14}}}section", name=name, id="{0}")
                    lst = etree.SubElement(section, f"{{{p14}}}sldIdLst")
                    for m in members:
                        etree.SubElement(lst, f"{{{p14}}}sldId", id=str(m))
                blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(item, blob)
    return out.getvalue()


def deck_at(tmp_path: Path, data: bytes) -> slides.Deck:
    path = tmp_path / "deck.pptx"
    path.write_bytes(data)
    return slides.read_deck(path)


def test_slides_are_read_from_the_decks_own_structure(tmp_path: Path) -> None:
    deck = deck_at(tmp_path, make_deck())
    assert deck.slide_count == 4 and not deck.sectioned
    s1, s2, s3, _ = deck.slides
    assert (s1.number, s1.title) == (1, "Continuous Integration")
    assert s1.items == ("Merge small changes often\n- Every push is built",)
    assert s1.notes == "Ask who has broken the build this week."
    table, chart = s2.items
    assert isinstance(table, slides.Table) and isinstance(chart, slides.Table)
    assert table.header == ("Stage", "Minutes") and table.rows == (("Build", "4"), ("Test", "11"))
    assert chart.caption == "Chart: Builds per day"
    assert chart.header == ("", "Builds") and chart.rows == (("Mon", "12"), ("Tue", "7.5"))
    assert s3.title is None
    [big, text] = s3.items  # the tiny picture is left out
    assert isinstance(big, slides.Picture) and text == "The pipeline runs on every push."
    assert big.context == "The pipeline runs on every push."


def test_paragraphs_put_each_slide_under_its_title(tmp_path: Path) -> None:
    deck = deck_at(tmp_path, make_deck())
    paras = slides.paragraphs(deck, {})
    headings = [(p.text, p.level, p.page) for p in paras if p.level is not None]
    assert headings == [("Continuous Integration", 0, 1), ("Build times", 0, 2), ("Secret slide", 0, 4)]
    notes = next(p for p in paras if p.text.startswith(slides.NOTES))
    assert notes.page == 1 and notes.path == ("Continuous Integration",)
    untitled = next(p for p in paras if p.page == 3)
    assert untitled.path == ("Build times",)
    tables = [p for p in paras if p.table]
    assert [p.page for p in tables] == [2, 2]


def test_sections_are_the_top_level(tmp_path: Path) -> None:
    deck = deck_at(tmp_path, make_deck(sections=True))
    assert deck.sectioned
    paras = slides.paragraphs(deck, {})
    headings = [(p.text, p.level) for p in paras if p.level is not None]
    assert headings == [
        ("Basics", 0), ("Continuous Integration", 1), ("Pipelines", 0), ("Build times", 1), ("Secret slide", 1),
    ]  # fmt: skip
    assert next(p for p in paras if p.text.startswith("Merge")).path == ("Basics", "Continuous Integration")


def test_hidden_slides_are_left_out_but_keep_their_numbers(tmp_path: Path) -> None:
    deck = deck_at(tmp_path, make_deck(hidden=True))
    assert deck.slide_count == 4 and [s.number for s in deck.slides] == [1, 2, 3]


def test_a_picture_on_most_slides_is_decoration(tmp_path: Path) -> None:
    deck = deck_at(tmp_path, make_deck(logo_everywhere=True))
    found = slides.pictures(deck)
    assert [len(p.blob) for p in found] == [len(picture())]


@pytest.mark.parametrize(
    ("data", "reason"),
    [
        (b"not a zip at all", "CORRUPT"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, "ENCRYPTED"),
    ],
    ids=["not-zip", "ole"],
)
def test_unreadable_decks_fail_with_their_reason(tmp_path: Path, data: bytes, reason: str) -> None:
    with pytest.raises(IngestError) as e:
        deck_at(tmp_path, data)
    assert e.value.reason == reason


def test_a_zip_that_unpacks_too_far_is_too_large(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(slides, "MAX_EXPANDED_BYTES", 1000)
    with pytest.raises(IngestError) as e:
        deck_at(tmp_path, make_deck())
    assert e.value.reason == "TOO_LARGE"


def test_a_deck_is_indexed_slide_by_slide(api: TestClient, embeddings: FakeOpenAI) -> None:
    doc = ingested(api, make_deck(), content_type=PPTX)
    assert doc["outcome"] == "indexed" and doc["page_count"] == 4 and doc["figures"] == []
    chunks = doc["chunks"]
    assert all(c["page_start"] == c["page_end"] for c in chunks)
    by_kind = {(c["page_start"], c["kind"]) for c in chunks}
    assert {(1, "text"), (2, "table"), (3, "text"), (4, "text")} <= by_kind
    first = chunks[0]
    assert first["section_path"] == "Continuous Integration"
    assert "Merge small changes often\n- Every push is built" in first["content"]
    assert "Speaker notes: Ask who has broken the build this week." in first["content"]
    chart = next(c for c in chunks if "Builds per day" in c["content"])
    assert "| Mon | 12 |" in chart["content"] and chart["section_path"] == "Build times"


def test_pictures_are_described_by_the_figure_model(
    api: TestClient, openai_fake: FakeOpenAI, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    with_vision(
        monkeypatch, VisionConfig("vision-model", tmp_path / "cache", 500, figure_model="figure-model")
    )
    useful = {"useful": True, "description": DESCRIPTION, "key_values": []}
    openai_fake.reply(json=chat_response(json.dumps(useful)))
    openai_fake.replies += [fake_embeddings]
    doc = ingested(api, make_deck(), content_type=PPTX)
    describe, _ = openai_fake.requests
    assert describe["model"] == "figure-model"
    assert "The pipeline runs on every push." in describe["messages"][0]["content"][0]["text"]
    [figure] = [c for c in doc["chunks"] if c["kind"] == "figure"]
    assert figure["page_start"] == 3 and figure["source"] == "vision" and DESCRIBED in figure["content"]
    assert DESCRIPTION in figure["content"] and figure["section_path"] == "Build times"
    [crop] = doc["figures"]
    assert crop["image"] == figure["image"]

    # a reindex calls no model: the description is cached by picture, the embeddings by text
    openai_fake.requests.clear()
    again = ingested(api, make_deck(), content_type=PPTX)
    assert openai_fake.requests == []
    assert [c["content"] for c in again["chunks"]] == [c["content"] for c in doc["chunks"]]


def test_a_picture_the_model_finds_useless_is_not_indexed(
    api: TestClient, openai_fake: FakeOpenAI, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    with_vision(
        monkeypatch, VisionConfig("vision-model", tmp_path / "cache", 500, figure_model="figure-model")
    )
    openai_fake.reply(json=chat_response(json.dumps({"useful": False, "description": "Clip art."})))
    openai_fake.replies += [fake_embeddings]
    doc = ingested(api, make_deck(), content_type=PPTX)
    assert doc["figures"] == [] and all(c["kind"] != "figure" for c in doc["chunks"])


def test_a_deck_with_too_many_slides_is_refused_before_any_model_call(
    api: TestClient, openai_fake: FakeOpenAI, monkeypatch: pytest.MonkeyPatch
) -> None:
    from conftest import configured

    monkeypatch.setattr("ingest.get_settings", lambda: configured(ingest_max_pages=3))
    doc = ingested(api, make_deck(), content_type=PPTX)
    assert (doc["outcome"], doc["failure_reason"]) == ("failed", "TOO_LARGE")
    assert openai_fake.requests == []
