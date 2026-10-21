"""PowerPoint (.pptx) → paragraphs, tables and figures, read from the deck's own structure. Nothing is guessed from
the layout as it is for a PDF: the file says which text is a slide's title, which shapes are tables and charts, and
where each slide begins.

A slide is a page: chunks cite its slide number, and ingest chunks a deck slide by slide (SLIDE_CHUNKS). Hidden
slides are left out but keep their numbers, so "slide 7" is the seventh slide the teacher sees in PowerPoint.

- Headings: each slide's title placeholder. A deck with two or more PowerPoint sections has its section names as
  the top level and slide titles under them. A slide without a title stays under the previous slide's, as a
  continued slide does in a PDF.
- Text: one paragraph per text box, one line per PowerPoint paragraph, sub-levels as indented "- " items.
  Shapes are read top to bottom, then left to right; a group's shapes in their own order at the group's place.
- Tables: a table, its first row the header when the table is styled with a header row.
- Charts: a table of the chart's own data (one row per category, one column per series), captioned with its
  title, so its numbers are exact and cost no model call.
- SmartArt: the text of its nodes, as lines.
- Pictures covering at least MIN_PICTURE_SHARE of the slide are described by the figure model, as a PDF's figures
  are (`describe`), with the slide's text as context, and indexed as figures; the same picture on REPEAT_SLIDES or
  more slides is decoration and is skipped. Pictures PyMuPDF can't read (EMF/WMF drawings) are skipped too.
- Speaker notes: a paragraph at the end of the slide, as notes often hold what the teacher says about it.

Size: the zip is checked before python-pptx expands it (MAX_EXPANDED_BYTES, MAX_PARTS), so a small upload
can't unpack into gigabytes.
"""

import hashlib
import logging
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pymupdf
from fastapi import HTTPException
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml import parse_xml

import pdf
from ingest.figures import Figure
from ingest.headings import Para, _region_para
from ingest.parse import IngestError
from ingest.prompts import DESCRIBE_PROMPT, DESCRIBE_PROMPT_VERSION
from ingest.strings import clean_cell
from ingest.tables import Table
from ingest.vision_calls import FIGURE_CONCURRENCY, _semaphore, ask, cache_file, cached

log = logging.getLogger("ai-core.ingest.slides")

MIN_PICTURE_SHARE = (
    0.05  # a picture smaller than this share of the slide isn't described (icons, logos, bullets)
)
REPEAT_SLIDES = 3  # the same picture on this many slides is decoration (a template's artwork)
MAX_PICTURE_SIDE = 2048  # pixels: bigger pictures are halved until they fit, before the model sees them
MAX_CONTEXT_WORDS = 300  # of the slide's text, given to the description call
MAX_EXPANDED_BYTES = 1 << 30  # the deck unpacked
MAX_PARTS = 20000  # files in the zip
NOTES = "Speaker notes:"
NO_BOX = (0.0, 0.0, 0.0, 0.0)
CFB_MAGIC = b"\xd0\xcf\x11\xe0"  # an OLE file: a password-protected deck, or a pre-2007 .ppt
SECTIONS = "http://schemas.microsoft.com/office/powerpoint/2010/main"  # PowerPoint 2010's sections
RELATIONSHIPS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}dm"


@dataclass(frozen=True)
class Picture:
    blob: bytes = field(repr=False)
    context: str  # the slide's title and text: what the picture is shown with


@dataclass(frozen=True)
class Slide:
    number: int  # 1-based, counting hidden slides
    title: str | None
    section: str | None
    items: tuple[str | Table | Picture, ...]  # text boxes, tables and pictures in reading order
    notes: str | None


@dataclass(frozen=True)
class Deck:
    slide_count: int
    slides: tuple[Slide, ...]  # the visible ones
    sectioned: bool  # two or more PowerPoint sections: their names are the top level


def read_deck(path: Path) -> Deck:
    """The deck's slides. CORRUPT, ENCRYPTED or TOO_LARGE when it can't be read."""
    _check_zip(path)
    try:
        return _read(path)
    except IngestError:
        raise
    except (
        Exception
    ) as e:  # a shape's own problem costs only that shape (_read_shape); this is the whole file
        raise IngestError("CORRUPT", f"python-pptx could not read the file: {e}") from None


def _read(path: Path) -> Deck:
    prs = Presentation(str(path))
    sections = _sections(prs)
    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        if slide._element.get("show") == "0":
            continue
        title_shape = slide.shapes.title
        title = _one_line(title_shape.text_frame.text) if title_shape is not None else ""
        items: list[str | Table | Picture] = []
        title_id = title_shape.shape_id if title_shape is not None else None
        for shape in _ordered(s for s in slide.shapes if s.shape_id != title_id):
            _read_shape(shape, slide.part, items)
        context = " ".join([title, *(i for i in items if isinstance(i, str))]).split()
        items = [
            Picture(i.blob, " ".join(context[:MAX_CONTEXT_WORDS])) if isinstance(i, Picture) else i
            for i in items
        ]
        notes = None
        if slide.has_notes_slide and (frame := slide.notes_slide.notes_text_frame) is not None:
            notes = frame.text.strip() or None
        slides.append(Slide(number, title or None, sections.get(slide.slide_id), tuple(items), notes))
    return Deck(len(prs.slides), tuple(slides), len(set(sections.values())) >= 2)


def pictures(deck: Deck) -> list[Picture]:
    """The pictures worth describing, each once: not repeated across REPEAT_SLIDES or more slides."""
    on_slides = Counter(
        h for s in deck.slides for h in {_hash(i.blob) for i in s.items if isinstance(i, Picture)}
    )
    seen: dict[str, Picture] = {}
    for s in deck.slides:
        for i in s.items:
            if isinstance(i, Picture) and on_slides[_hash(i.blob)] < REPEAT_SLIDES:
                seen.setdefault(_hash(i.blob) + _hash(i.context.encode()), i)
    return list(seen.values())


async def describe(
    found: list[Picture], model: str | None, cache_dir: Path, max_pictures: int
) -> dict[Picture, Figure]:
    """Figures for the pictures the model found useful, by picture. None (figures off): no calls, no figures."""
    if model is None:
        return {}
    out: dict[Picture, Figure] = {}
    for picture in found[:max_pictures]:
        if (figure := await _described(picture, model, cache_dir)) is not None:
            out[picture] = figure
    return out


def paragraphs(deck: Deck, figures: dict[Picture, Figure]) -> list[Para]:
    """The deck as paragraphs, headings, tables and figures for the chunker, slide by slide."""
    paras: list[Para] = []
    stack: list[tuple[int, str]] = []
    section: str | None = None
    title_level = 1 if deck.sectioned else 0
    for slide in deck.slides:
        page = slide.number
        if deck.sectioned and slide.section and slide.section != section:
            section = slide.section
            stack = [(0, section)]
            paras.append(Para(section, page, (section,), 0))
        if slide.title:
            while stack and stack[-1][0] >= title_level:
                stack.pop()
            stack.append((title_level, slide.title))
            paras.append(Para(slide.title, page, tuple(t for _, t in stack), title_level))
        for item in slide.items:
            if isinstance(item, str):
                paras.append(Para(item, page, tuple(t for _, t in stack)))
            elif isinstance(item, Table):
                paras.append(_region_para(item, page, stack, paras))
            elif item in figures:
                paras.append(_region_para(figures[item], page, stack, paras))
        if slide.notes:
            paras.append(Para(f"{NOTES} {slide.notes}", page, tuple(t for _, t in stack)))
    return paras


def _check_zip(path: Path) -> None:
    with path.open("rb") as f:
        if f.read(len(CFB_MAGIC)) == CFB_MAGIC:
            raise IngestError(
                "ENCRYPTED", "The file is password protected, or an old .ppt rather than a .pptx"
            )
    try:
        with zipfile.ZipFile(path) as z:
            entries = z.infolist()
    except (zipfile.BadZipFile, OSError) as e:
        raise IngestError("CORRUPT", f"Not a .pptx (zip) file: {e}") from None
    if len(entries) > MAX_PARTS or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
        raise IngestError("TOO_LARGE", f"The deck unpacks to more than {MAX_EXPANDED_BYTES >> 20} MiB")


def _sections(prs: Any) -> dict[int, str]:
    """Slide id → the name of its PowerPoint section."""
    out: dict[int, str] = {}
    for section in prs.part._element.xpath(f".//*[local-name()='section' and namespace-uri()='{SECTIONS}']"):
        name = _one_line(section.get("name") or "")
        for slide in section.xpath(".//*[local-name()='sldId']"):
            if name and (ident := slide.get("id", "")).isdigit():
                out[int(ident)] = name
    return out


def _ordered(shapes: Any) -> list[Any]:
    return sorted(shapes, key=lambda s: (s.top or 0, s.left or 0))


def _read_shape(shape: Any, part: Any, items: list[str | Table | Picture]) -> None:
    """Add what the shape holds to `items`. A shape python-pptx can't read is skipped, not the deck."""
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in _ordered(shape.shapes):
                _read_shape(child, part, items)
        elif shape.has_table:
            items.extend(_table(shape.table))
        elif shape.has_chart:
            items.extend(_chart(shape.chart))
        elif _is_picture(shape):
            if shape.width and shape.height and _share(shape) >= MIN_PICTURE_SHARE:
                items.append(Picture(shape.image.blob, ""))
        elif shape.has_text_frame:
            if text := _frame_text(shape.text_frame):
                items.append(text)
        elif text := _smartart(shape, part):
            items.append(text)
    except Exception as e:  # a malformed shape costs only itself
        log.warning("shape_skipped", extra={"shape": getattr(shape, "name", "?"), "error": str(e)})


def _share(shape: Any) -> float:
    slide = shape.part.package.presentation_part.presentation
    return (shape.width * shape.height) / max(1, (slide.slide_width or 1) * (slide.slide_height or 1))


def _is_picture(shape: Any) -> bool:
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return True
    try:
        return bool(shape.is_placeholder and shape.image)  # a picture placeholder holding an image
    except (AttributeError, ValueError):
        return False


def _frame_text(frame: Any) -> str:
    lines = []
    for p in frame.paragraphs:
        text = _one_line("".join(r.text for r in p.runs) if p.runs else p.text)
        if text:
            lines.append(f"{'  ' * (p.level - 1)}- {text}" if p.level else text)
    return "\n".join(lines)


def _table(table: Any) -> list[str | Table]:
    rows = [tuple("" if c.is_spanned else clean_cell(c.text) for c in row.cells) for row in table.rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    header = rows[0] if table.first_row and len(rows) > 1 else None
    body = rows[1:] if header else rows
    return [Table(NO_BOX, header, tuple(body))]


def _chart(chart: Any) -> list[str | Table]:
    caption = _one_line(chart.chart_title.text_frame.text) if chart.has_title else None
    if not chart.plots:
        return [caption] if caption else []
    plot = chart.plots[0]
    series = list(plot.series)
    categories = [clean_cell(c) for c in plot.categories]
    if not series or not categories:
        return [caption] if caption else []
    header = ("", *(clean_cell(s.name) for s in series))
    values = [list(s.values) for s in series]
    rows = tuple(
        (category, *(_number(v[i]) if i < len(v) else "" for v in values))
        for i, category in enumerate(categories)
    )
    return [Table(NO_BOX, header, rows, f"Chart: {caption}" if caption else "Chart")]


def _number(value: object) -> str:
    if value is None:
        return ""
    return f"{value:g}" if isinstance(value, float) else str(value)


def _smartart(shape: Any, part: Any) -> str:
    """A SmartArt graphic's node texts, one per line, or "" for any other shape."""
    ids = shape._element.xpath(".//*[local-name()='relIds']")
    if not ids or not (rid := ids[0].get(RELATIONSHIPS)):
        return ""
    data = parse_xml(part.related_part(rid).blob)
    lines = []
    for point in data.xpath(".//*[local-name()='pt']"):
        if point.get("type") not in (None, "node"):
            continue  # the document node and the connectors between nodes
        if text := _one_line(" ".join(point.xpath(".//*[local-name()='t']/text()"))):
            lines.append(f"- {text}")
    return "\n".join(lines)


def _one_line(text: str) -> str:
    return " ".join(text.replace("\x0b", " ").split())


def _hash(blob: bytes) -> str:
    return hashlib.sha256(blob).hexdigest()


def _png(blob: bytes) -> bytes | None:
    """The picture as a PNG the model can read and the browser can show, or None for a format PyMuPDF can't read."""
    try:
        pix = pymupdf.Pixmap(blob)
        if pix.alpha:
            pix = pymupdf.Pixmap(pix, 0)
        if pix.colorspace is None or pix.colorspace.n not in (1, 3):
            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        while max(pix.w, pix.h) > MAX_PICTURE_SIDE:
            pix.shrink(1)
        return pix.tobytes("png")
    except Exception:
        return None


async def _described(picture: Picture, model: str, cache_dir: Path) -> Figure | None:
    png = await pdf.in_thread(_png, picture.blob)
    if png is None:
        return None
    prompt = DESCRIBE_PROMPT.format(
        type="picture on a presentation slide",
        caption="",
        words="(none)",
        context=picture.context or "(none)",
    )
    what = _hash(picture.context.encode())[:8]
    try:
        raw = await cached(
            cache_file(cache_dir, model, _hash(png), f"pptx-d{DESCRIBE_PROMPT_VERSION}-{what}"),
            lambda: ask(prompt, png, model, _semaphore("figure", FIGURE_CONCURRENCY)),
            "describe",
        )
    except HTTPException:
        raise  # the teacher's allowance ran out: the ingest stops
    except Exception as e:
        log.warning("picture_not_described", extra={"error": str(e)})
        return None
    description = " ".join(str(raw.get("description") or "").split())
    if raw.get("useful") is False or not description:
        return None
    values = tuple(
        v
        for v in (
            " ".join(str(x).split()) for x in raw.get("key_values") or [] if isinstance(x, str | int | float)
        )
        if v
    )
    return Figure(NO_BOX, "picture", description, None, values, "", png)
